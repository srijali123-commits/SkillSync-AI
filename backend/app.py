from flask import Flask, request, jsonify
from flask_cors import CORS
from datetime import datetime, timezone
import io
import os
import re
import time
import threading
import copy
from pathlib import Path
import requests
from dotenv import load_dotenv

load_dotenv(dotenv_path=Path(__file__).resolve().parent / '.env')

try:
    from PyPDF2 import PdfReader
except ImportError:
    PdfReader = None
try:
    from docx import Document
except ImportError:
    Document = None

app = Flask(__name__)
CORS(app)
app.config['MAX_CONTENT_LENGTH'] = 5 * 1024 * 1024
GITHUB_API = 'https://api.github.com'
LEETCODE_API = 'https://leetcode.com/graphql/'
REQUEST_TIMEOUT = 15

ROLE_SKILLS = {
    'Software Engineer': ['Programming', 'Data Structures', 'Algorithms', 'Git', 'Problem Solving'],
    'Frontend Developer': ['HTML', 'CSS', 'JavaScript', 'React', 'Git'],
    'Backend Developer': ['Python', 'Java', 'APIs', 'Databases', 'Git'],
    'Full Stack Developer': ['JavaScript', 'React', 'Node.js', 'Databases', 'APIs'],
    'AI/ML Engineer': ['Python', 'Machine Learning', 'Data Structures', 'NumPy', 'TensorFlow'],
    'Data Scientist': ['Python', 'SQL', 'Statistics', 'Pandas', 'Machine Learning'],
    'DevOps Engineer': ['Linux', 'Git', 'Docker', 'Cloud', 'CI/CD'],
    'Android Developer': ['Java', 'Kotlin', 'Android Studio', 'APIs', 'Git'],
}
SKILL_ALIASES = {
    'html': 'HTML', 'html5': 'HTML', 'css': 'CSS', 'css3': 'CSS',
    'javascript': 'JavaScript', 'react': 'React', 'reactjs': 'React', 'react.js': 'React',
    'python': 'Python', 'java': 'Java', 'node': 'Node.js', 'nodejs': 'Node.js',
    'node.js': 'Node.js', 'sql': 'SQL', 'database': 'Databases',
    'databases': 'Databases', 'mysql': 'Databases', 'mongodb': 'Databases',
    'postgresql': 'Databases', 'dbms': 'Databases', 'api': 'APIs',
    'apis': 'APIs', 'rest api': 'APIs', 'rest apis': 'APIs',
    'restful api': 'APIs', 'git': 'Git', 'github': 'Git', 'gitlab': 'Git',
    'docker': 'Docker', 'linux': 'Linux', 'cloud': 'Cloud', 'aws': 'Cloud',
    'azure': 'Cloud', 'gcp': 'Cloud', 'machine learning': 'Machine Learning',
    'machine-learning': 'Machine Learning', 'ml': 'Machine Learning',
    'tensorflow': 'TensorFlow', 'numpy': 'NumPy', 'pandas': 'Pandas',
    'statistics': 'Statistics', 'statistical analysis': 'Statistics',
    'kotlin': 'Kotlin', 'android studio': 'Android Studio',
    'data structures': 'Data Structures', 'data structure': 'Data Structures',
    'dsa': 'Data Structures', 'data structures & algorithms': 'Data Structures',
    'data structures and algorithms': 'Data Structures', 'algorithms': 'Algorithms',
    'algorithm': 'Algorithms', 'problem solving': 'Problem Solving',
    'problem-solving': 'Problem Solving', 'programming': 'Programming',
    'ci/cd': 'CI/CD', 'cicd': 'CI/CD',
}

def normalize_text(text):
    text = str(text or '').replace('\x00', ' ')
    text = re.sub(r'\r\n?', '\n', text)
    text = re.sub(r'[ \t]+', ' ', text)
    return re.sub(r'\n{3,}', '\n\n', text).strip()

def safe_int(value, default=0):
    try:
        return int(value)
    except (ValueError, TypeError, OverflowError):
        return default

def safe_float(value, default=0):
    try:
        result = float(value)
        return result if result == result and result not in (float('inf'), float('-inf')) else default
    except (ValueError, TypeError, OverflowError):
        return default

def clamp(value, low=0, high=100):
    return max(low, min(high, safe_float(value)))

def percentage(part, total):
    return round(part / total * 100, 1) if total else 0

def contains_term(text, term):
    return re.search(r'(?<![a-zA-Z0-9])' + re.escape(term) + r'(?![a-zA-Z0-9])', text or '', re.I) is not None

def detect_skills(text):
    detected = []
    for alias, skill in sorted(SKILL_ALIASES.items(), key=lambda item: len(item[0]), reverse=True):
        if contains_term(text, alias) and skill not in detected:
            detected.append(skill)
    if re.search(r'data structures?\s*(?:&|and)\s*algorithms?', text or '', re.I):
        for skill in ('Data Structures', 'Algorithms'):
            if skill not in detected:
                detected.append(skill)
    return detected

def now_iso():
    return datetime.now(timezone.utc).isoformat()

def normalize_username(value, platform):
    username = str(value or '').strip().rstrip('/')
    if '/' in username:
        username = username.split('/')[-1]
    if platform == 'github':
        valid = re.fullmatch(r'[A-Za-z0-9](?:[A-Za-z0-9-]{0,37}[A-Za-z0-9])?', username)
        if not valid or '--' in username:
            return ''
    elif not re.fullmatch(r'[A-Za-z0-9_-]{1,50}', username):
        return ''
    return username

def first_value(data, *keys, default=None):
    for key in keys:
        if key in data and data[key] is not None:
            return data[key]
    return default

def as_bool(value, default=False):
    if value is None:
        return default
    if isinstance(value, bool):
        return value
    if isinstance(value, str):
        return value.strip().lower() in ('true', '1', 'yes', 'on')
    return bool(value)

def valid_role(value):
    return value if value in ROLE_SKILLS else 'Software Engineer'

# RESUME EXTRACTION

def extract_resume_text(file):
    filename = (file.filename or '').lower()
    raw = file.read()
    try:
        if filename.endswith('.pdf'):
            if PdfReader is None:
                return ''
            reader = PdfReader(io.BytesIO(raw))
            return normalize_text('\n'.join(page.extract_text() or '' for page in reader.pages))
        if filename.endswith('.docx'):
            if Document is None:
                return ''
            doc = Document(io.BytesIO(raw))
            paragraphs = [paragraph.text for paragraph in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        paragraphs.append(cell.text)
            return normalize_text('\n'.join(paragraphs))
        if filename.endswith('.txt'):
            return normalize_text(raw.decode('utf-8', errors='ignore'))
    except Exception as error:
        app.logger.warning('Resume extraction failed: %s', error)
    return ''

SECTION_HEADINGS = [
    'Career Objective', 'Objective', 'Summary', 'Profile', 'Professional Summary',
    'Education', 'Academic Background', 'Educational Qualification',
    'Technical Skills', 'Technical Skill', 'Skills', 'Technology Stack',
    'Projects', 'Project', 'Academic Projects', 'Personal Projects',
    'Certifications', 'Certification', 'Certificates', 'Experience',
    'Work Experience', 'Professional Experience', 'Internship', 'Internships',
    'Internship Experience', 'Internship Training', 'Training & Internship',
    'Work & Internship Experience', 'Industry Experience', 'Achievements',
    'Soft Skills', 'Declaration', 'Contact',
]

def clean_heading(line):
    return re.sub(r'^[#•*\-:\s]+|[:\s]+$', '', str(line or '')).strip().lower()

def find_section(text, headings):
    lines = [line.strip() for line in normalize_text(text).splitlines() if line.strip()]
    requested = {heading.lower() for heading in headings}
    all_headings = {heading.lower() for heading in SECTION_HEADINGS}
    start = next((i + 1 for i, line in enumerate(lines) if clean_heading(line) in requested), None)
    if start is None:
        return ''
    result = []
    for line in lines[start:]:
        if clean_heading(line) in all_headings:
            break
        result.append(line)
    return '\n'.join(result).strip()

def score_resume_quality(resume_text):
    text = normalize_text(resume_text)
    lower = text.lower()
    breakdown = []
    def add(label, score, maximum, details):
        breakdown.append({'category': label, 'label': label, 'score': round(score, 1),
                          'max': maximum, 'max_score': maximum,
                          'status': 'Present' if score > 0 else 'Missing', 'details': details})
    education = find_section(text, ['Education', 'Academic Background', 'Educational Qualification'])
    education_found = bool(education) or any(word in lower for word in ['b.tech', 'btech', 'bachelor', 'university', 'college', 'cgpa'])
    add('Education', 2 if education_found else 0, 2, 'Education-related information was detected.' if education_found else 'Education information was not detected.')
    skills_section = find_section(text, ['Technical Skills', 'Technical Skill', 'Skills', 'Technology Stack'])
    skills = detect_skills(skills_section or text)
    technical_score = min(4, round(len(skills) / 3, 1))
    add('Technical Skills', technical_score, 4, f'{len(skills)} recognizable technical skills detected.')
    projects = find_section(text, ['Projects', 'Project', 'Academic Projects', 'Personal Projects'])
    signals = ['developed', 'built', 'implemented', 'created', 'designed', 'python', 'java', 'javascript', 'html', 'css', 'ai', 'machine learning', 'application', 'platform', 'api', 'database']
    hits = sum(signal in projects.lower() for signal in signals)
    project_score = min(5, round(2 + hits * 0.3, 1)) if projects else 0
    add('Projects', project_score, 5, 'Project evidence detected.' if projects else 'No dedicated projects section detected.')
    certifications = find_section(text, ['Certifications', 'Certification', 'Certificates'])
    cert_found = bool(certifications) or any(word in lower for word in ['certification', 'certificate', 'nptel', 'coursera', 'udemy'])
    add('Certifications', 2 if cert_found else 0, 2, 'Certification information detected.' if cert_found else 'Certifications not detected.')
    experience = find_section(text, ['Experience', 'Work Experience', 'Professional Experience', 'Internship', 'Internships', 'Internship Experience', 'Internship Training', 'Training & Internship', 'Work & Internship Experience', 'Industry Experience'])
    experience_found = bool(experience)
    add('Experience/Internship', 4 if experience_found else 0, 4, 'Experience or internship section detected.' if experience_found else 'No dedicated experience or internship section detected.')
    contact_found = bool(re.search(r'[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}', text)) or 'linkedin.com' in lower
    summary_found = bool(find_section(text, ['Career Objective', 'Objective', 'Summary', 'Profile', 'Professional Summary']))
    completeness_score = min(3, sum([education_found, bool(skills_section), bool(projects), contact_found]))
    add('Resume Completeness', completeness_score, 3, 'Completeness based on major resume sections and contact details.')
    total_score = round(sum(item['score'] for item in breakdown), 1)
    sections = {'education': education_found, 'technical_skills': technical_score > 0,
                'skills': technical_score > 0, 'projects': project_score > 0,
                'certifications': cert_found, 'experience': experience_found,
                'internship': experience_found, 'resume_completeness': completeness_score > 0,
                'contact': contact_found, 'summary': summary_found, 'objective': summary_found}
    return {'score': total_score, 'max': 20, 'max_score': 20, 'sections': sections,
            'breakdown': breakdown, 'detected_skills': detect_skills(text),
            'suggestions': [f"Improve {item['label']}." for item in breakdown if item['score'] < item['max']]}

# GITHUB CACHE
GITHUB_CACHE = {}
GITHUB_CACHE_LOCK = threading.Lock()
GITHUB_CACHE_TTL = 30 * 60

def get_cached_github(username, allow_expired=False):
    with GITHUB_CACHE_LOCK:
        entry = GITHUB_CACHE.get(str(username).lower())
        if not entry:
            return None
        if time.time() - entry['timestamp'] <= GITHUB_CACHE_TTL or allow_expired:
            return copy.deepcopy(entry['data'])
    return None

def save_cached_github(username, data):
    if data.get('available'):
        with GITHUB_CACHE_LOCK:
            GITHUB_CACHE[str(username).lower()] = {'timestamp': time.time(), 'data': copy.deepcopy(data)}

# GITHUB API

def github_headers():
    headers = {'Accept': 'application/vnd.github+json', 'User-Agent': 'SkillSync-AI',
               'X-GitHub-Api-Version': '2022-11-28'}
    token = os.getenv('GITHUB_TOKEN', '').strip()
    if token:
        headers['Authorization'] = f'Bearer {token}'
    return headers

def github_request(path, params=None):
    response = requests.get(GITHUB_API + path, headers=github_headers(), params=params, timeout=REQUEST_TIMEOUT)
    if response.status_code == 404:
        raise ValueError('GitHub profile not found.')
    if response.status_code == 401:
        raise ValueError('GitHub token was rejected (HTTP 401). Check whether GITHUB_TOKEN is valid and unexpired.')
    if response.status_code in (403, 429):
        remaining = response.headers.get('X-RateLimit-Remaining')
        reset = response.headers.get('X-RateLimit-Reset')
        if remaining == '0':
            reset_time = ''
            if reset and reset.isdigit():
                reset_time = ' Resets at ' + datetime.fromtimestamp(int(reset), timezone.utc).strftime('%Y-%m-%d %H:%M UTC') + '.'
            raise ValueError('GitHub API primary rate limit reached.' + reset_time)
        raise ValueError('GitHub API request restricted (HTTP %s). Possible secondary rate limit or permissions issue.' % response.status_code)
    if response.status_code != 200:
        raise ValueError(f'GitHub API returned HTTP {response.status_code}.')
    return response.json()

def github_unavailable(username, error):
    return {'available': False, 'score': 0, 'max_score': 20, 'username': username,
            'repositories': [], 'languages': [], 'score_breakdown': [], 'error': error, 'message': error}

def analyze_github(username, role='Software Engineer'):
    original = str(username or '').strip()
    if not original:
        return github_unavailable('', 'GitHub username was not provided.')
    username = normalize_username(original, 'github')
    if not username:
        return github_unavailable(original, 'Invalid GitHub username format.')
    cached = get_cached_github(username)
    if cached:
        cached['cached'] = True
        return cached
    try:
        profile = github_request(f'/users/{username}')
        repositories = github_request(f'/users/{username}/repos', {'per_page': 100, 'sort': 'updated', 'type': 'owner'})
        if not isinstance(repositories, list):
            raise ValueError('Unexpected GitHub repository response.')
        public_repos = safe_int(profile.get('public_repos'))
        followers = safe_int(profile.get('followers'))
        following = safe_int(profile.get('following'))
        active_repos = [repo for repo in repositories if isinstance(repo, dict) and not repo.get('fork')]
        languages = sorted({repo.get('language') for repo in repositories if isinstance(repo, dict) and repo.get('language')})
        repo_points = min(4, public_repos / 5)
        community_points = min(2, followers / 10)
        activity_points = min(5, len(active_repos) / 4)
        profile_points = sum([bool(profile.get('name')), bool(profile.get('bio')), bool(profile.get('avatar_url'))])
        technology_points = min(3, len(languages) / 2)
        quality_points = 0
        for repo in active_repos:
            if repo.get('description'):
                quality_points += 0.5
            if safe_int(repo.get('stargazers_count')) > 0:
                quality_points += 0.5
        quality_points = min(3, quality_points)
        breakdown = [('Public repositories', repo_points, 4), ('Community presence', community_points, 2),
                     ('Development activity', activity_points, 5), ('Profile completeness', profile_points, 3),
                     ('Technology diversity', technology_points, 3), ('Repository quality', quality_points, 3)]
        score = round(min(20, sum(item[1] for item in breakdown)), 1)
        result = {'available': True, 'score': score, 'max_score': 20, 'username': username,
                  'name': profile.get('name') or '', 'bio': profile.get('bio') or '',
                  'avatar_url': profile.get('avatar_url') or '', 'public_repos': public_repos,
                  'followers': followers, 'following': following, 'languages': languages,
                  'repositories': repositories[:20], 'active_repositories': len(active_repos),
                  'profile_completeness': profile_points,
                  'activity_signals': {'public_repositories': public_repos, 'followers': followers, 'following': following},
                  'role_evidence': [],
                  'score_breakdown': [{'label': label, 'score': round(points, 1), 'max': maximum,
                                       'max_score': maximum} for label, points, maximum in breakdown],
                  'error': None, 'cached': False}
        save_cached_github(username, result)
        return result
    except (requests.RequestException, ValueError) as error:
        app.logger.warning('GitHub lookup failed for %s: %s', username, error)
        previous = get_cached_github(username, allow_expired=True)
        if previous:
            previous['cached'] = True
            previous['stale'] = True
            previous['warning'] = 'GitHub API temporarily unavailable. Showing previously verified statistics.'
            return previous
        return github_unavailable(username, str(error))

# LEETCODE
LEETCODE_QUERY = '''
query userProfile($username: String!) {
  matchedUser(username: $username) {
    username
    profile { ranking }
    submitStatsGlobal { acSubmissionNum { difficulty count } }
  }
}
'''

def leetcode_unavailable(username, error):
    return {'available': False, 'score': 0, 'max_score': 20, 'username': username,
            'total_solved': 0, 'easy_solved': 0, 'medium_solved': 0, 'hard_solved': 0,
            'ranking': None, 'score_breakdown': [], 'error': error, 'message': error}

def leetcode_score(easy, medium, hard):
    total = easy + medium + hard
    easy_points = min(5, easy / 20)
    medium_points = min(7, medium / 15)
    hard_points = min(5, hard / 8)
    depth = min(3, ((medium + hard) / total) * 1.5 if total else 0)
    score = round(min(20, easy_points + medium_points + hard_points + depth), 1)
    breakdown = [('Easy problems', easy_points, 5), ('Medium problems', medium_points, 7),
                 ('Hard problems', hard_points, 5), ('Problem-solving depth', depth, 3)]
    return score, breakdown

def analyze_leetcode(username):
    original = str(username or '').strip()
    if not original:
        return leetcode_unavailable('', 'LeetCode username was not provided.')
    username = normalize_username(original, 'leetcode')
    if not username:
        return leetcode_unavailable(original, 'Invalid LeetCode username format.')
    try:
        response = requests.post(LEETCODE_API,
            json={'query': LEETCODE_QUERY, 'variables': {'username': username}},
            headers={'Content-Type': 'application/json', 'Accept': 'application/json',
                     'Origin': 'https://leetcode.com', 'Referer': f'https://leetcode.com/u/{username}/',
                     'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/120.0 Safari/537.36'},
            timeout=20)
        if response.status_code in (403, 429):
            raise ValueError('LeetCode blocked or rate-limited the API request.')
        if response.status_code != 200:
            raise ValueError(f'LeetCode API returned HTTP {response.status_code}.')
        try:
            payload = response.json()
        except ValueError:
            raise ValueError('LeetCode returned a non-JSON response.')
        if payload.get('errors'):
            raise ValueError('LeetCode GraphQL returned an error.')
        user = (payload.get('data') or {}).get('matchedUser')
        if not user:
            return leetcode_unavailable(username, 'LeetCode user not found.')
        stats = user.get('submitStatsGlobal')
        if not isinstance(stats, dict):
            raise ValueError('LeetCode did not return submission statistics.')
        entries = stats.get('acSubmissionNum')
        if not isinstance(entries, list):
            raise ValueError('LeetCode problem counts are unavailable.')
        counts = {str(item.get('difficulty')): safe_int(item.get('count')) for item in entries if isinstance(item, dict)}
        easy, medium, hard = counts.get('Easy', 0), counts.get('Medium', 0), counts.get('Hard', 0)
        total = counts.get('All', easy + medium + hard)
        ranking = (user.get('profile') or {}).get('ranking')
        score, breakdown = leetcode_score(easy, medium, hard)
        return {'available': True, 'score': score, 'max_score': 20,
                'username': user.get('username') or username, 'total_solved': total,
                'easy_solved': easy, 'medium_solved': medium, 'hard_solved': hard,
                'ranking': ranking,
                'difficulty_depth': {'easy': easy, 'medium': medium, 'hard': hard,
                                     'medium_hard_percentage': percentage(medium + hard, total)},
                'score_breakdown': [{'label': label, 'score': round(points, 1), 'max': maximum,
                                     'max_score': maximum} for label, points, maximum in breakdown],
                'error': None}
    except (requests.RequestException, ValueError) as error:
        app.logger.warning('LeetCode lookup failed for %s: %s', username, error)
        return leetcode_unavailable(username, str(error))

# ROLE MATCHING

def role_skill_analysis(resume_text, role):
    role = valid_role(role)
    required = ROLE_SKILLS[role]
    detected = detect_skills(resume_text)
    matched = [skill for skill in required if skill in detected]
    missing = [skill for skill in required if skill not in detected]
    return {'role': role, 'required_skills': required, 'detected_skills': detected,
            'matched_skills': matched, 'missing_skills': missing,
            'match_percentage': percentage(len(matched), len(required))}

def build_skill_gap(role, resume_text, github_data=None, leetcode_data=None):
    result = role_skill_analysis(resume_text, role)
    gaps = []
    for skill in result['required_skills']:
        missing = skill in result['missing_skills']
        if not missing:
            priority = 'Covered'
        elif skill in ('Data Structures', 'Algorithms', 'Problem Solving'):
            priority = 'Critical'
        elif skill in ('Programming', 'Git', 'APIs', 'Databases', 'Machine Learning'):
            priority = 'Important'
        else:
            priority = 'Optional'
        gaps.append({'skill': skill, 'priority': priority, 'status': 'Missing' if missing else 'Detected',
                     'evidence': 'No clear evidence found in resume.' if missing else 'Skill detected in resume.',
                     'action': f'Build a project or practice {skill}.' if missing else 'Maintain and strengthen this skill.'})
    return {**result, 'gaps': gaps}

# CAREER READINESS

def calculate_readiness(resume_score, github_score, github_available, leetcode_score,
                        leetcode_available, role_match_percentage):
    role_component = round(clamp(role_match_percentage) * 0.40, 1)
    resume_component = clamp(resume_score, 0, 20)
    github_component = clamp(github_score, 0, 20) if github_available else 0
    leetcode_component = clamp(leetcode_score, 0, 20) if leetcode_available else 0
    overall = round(role_component + resume_component + github_component + leetcode_component, 1)
    coverage = 60 + (20 if github_available else 0) + (20 if leetcode_available else 0)
    level = ('Excellent' if overall >= 85 else 'Strong' if overall >= 70 else
             'Developing' if overall >= 55 else 'Needs Improvement' if overall > 0 else 'Not Available')
    components = [
        {'name': 'Role-Specific Skills', 'label': 'Role-Specific Skills', 'score': role_component,
         'max': 40, 'max_score': 40, 'available': True},
        {'name': 'GitHub Profile & Development Activity', 'label': 'GitHub Profile & Development Activity',
         'score': github_component, 'max': 20, 'max_score': 20, 'available': bool(github_available)},
        {'name': 'LeetCode Problem-Solving', 'label': 'LeetCode Problem-Solving',
         'score': leetcode_component, 'max': 20, 'max_score': 20, 'available': bool(leetcode_available)},
        {'name': 'Resume/Profile Quality', 'label': 'Resume/Profile Quality',
         'score': resume_component, 'max': 20, 'max_score': 20, 'available': True},
    ]
    return {'overall_score': overall, 'score': overall, 'level': level, 'confidence': coverage,
            'data_coverage': coverage, 'available_weight': coverage, 'components': components,
            'formula': 'Role Skills 40% + GitHub 20% + LeetCode 20% + Resume Quality 20%'}

def run_analysis(resume_text, role, github_username=None, leetcode_username=None):
    resume_text = normalize_text(resume_text)
    role = valid_role(role)
    resume_data = score_resume_quality(resume_text)
    role_data = role_skill_analysis(resume_text, role)
    github_data = analyze_github(github_username, role)
    leetcode_data = analyze_leetcode(leetcode_username)
    readiness = calculate_readiness(resume_score=resume_data['score'],
        github_score=github_data['score'], github_available=github_data['available'],
        leetcode_score=leetcode_data['score'], leetcode_available=leetcode_data['available'],
        role_match_percentage=role_data['match_percentage'])
    skill_gap = build_skill_gap(role, resume_text, github_data, leetcode_data)
    return {'timestamp': now_iso(), 'target_role': role, 'resume_text': resume_text,
            'resume': resume_data, 'github': github_data, 'leetcode': leetcode_data,
            'role_matching': role_data, 'skill_gap': skill_gap,
            'required_skills': role_data['required_skills'], 'detected_skills': role_data['detected_skills'],
            'missing_skills': role_data['missing_skills'], 'readiness': readiness,
            'readiness_score': readiness['overall_score'], 'readiness_level': readiness['level'],
            'confidence': readiness['confidence'], 'data_coverage': readiness['data_coverage']}

# RECOMMENDATIONS

def generate_week1_recommendations(analysis):
    recommendations = []
    for gap in (analysis.get('skill_gap') or {}).get('gaps', []):
        if gap.get('status') == 'Missing':
            recommendations.append({'title': f"Improve {gap['skill']}", 'skill': gap['skill'],
                                    'priority': gap['priority'], 'reason': gap['evidence'],
                                    'action': gap['action']})
    for suggestion in (analysis.get('resume') or {}).get('suggestions', []):
        recommendations.append({'title': 'Improve resume quality', 'skill': 'Resume',
                                'priority': 'Important', 'reason': suggestion, 'action': suggestion})
    return recommendations[:8]

# WHAT-IF SIMULATOR

def simulate_what_if(analysis, scenario):
    scenario = normalize_text(scenario)
    if not scenario:
        raise ValueError('Please enter a career scenario.')
    role = valid_role(analysis.get('target_role', 'Software Engineer'))
    role_data = analysis.get('role_matching') or {}
    required = role_data.get('required_skills') or ROLE_SKILLS[role]
    detected = set(role_data.get('detected_skills') or [])
    current = safe_float((analysis.get('readiness') or {}).get('overall_score', analysis.get('readiness_score', 0)))
    projected = current
    explanations, assumptions = [], []
    if re.search(r'\b(learn|master|study|acquire|gain|build|practice)\b', scenario, re.I):
        new_skills = [skill for skill in required if skill in detect_skills(scenario) and skill not in detected]
        if new_skills and required:
            gain = round(40 * len(new_skills) / len(required), 1)
            projected += gain
            explanations.append('New role skills: ' + ', '.join(new_skills) + f' (+{gain}).')
    match = re.search(r'\b(?:solve|solved|solving|complete|completed|practice|practiced)\s+(\d+)\s+(?:(easy|medium|hard)\s+)?(?:leetcode\s+)?(?:questions?|problems?)\b', scenario, re.I)
    if not match:
        match = re.search(r'\b(\d+)\s+(?:(easy|medium|hard)\s+)?leetcode\s+(?:questions?|problems?)\b', scenario, re.I)
    if match:
        count = min(safe_int(match.group(1)), 10000)
        difficulty = (match.group(2) or '').lower()
        lc = analysis.get('leetcode') or {}
        available = bool(lc.get('available'))
        easy = safe_int(lc.get('easy_solved')) if available else 0
        medium = safe_int(lc.get('medium_solved')) if available else 0
        hard = safe_int(lc.get('hard_solved')) if available else 0
        previous, _ = leetcode_score(easy, medium, hard)
        if difficulty == 'easy':
            easy += count
        elif difficulty == 'medium':
            medium += count
        elif difficulty == 'hard':
            hard += count
        else:
            easy_added = round(count * 0.5)
            medium_added = round(count * 0.4)
            easy += easy_added
            medium += medium_added
            hard += count - easy_added - medium_added
            assumptions.append('Assumed 50% easy, 40% medium, 10% hard.')
        if not available:
            assumptions.append('LeetCode data unavailable; zero baseline assumed.')
        updated, _ = leetcode_score(easy, medium, hard)
        gain = max(0, round(updated - previous, 1))
        projected += gain
        explanations.append(f'LeetCode score: {previous}/20 → {updated}/20 (+{gain}).')
    projected = round(min(100, projected), 1)
    improvement = round(projected - current, 1)
    message = (' '.join(explanations) if explanations else
               'No measurable component change was identified. Specify a new role skill or number of LeetCode questions.')
    if assumptions:
        message += ' Assumptions: ' + ' '.join(assumptions)
    return {'scenario': scenario, 'target_role': role, 'current_score': current,
            'projected_score': projected, 'improvement': improvement, 'message': message,
            'assumptions': assumptions, 'estimated': True}

# API ROUTES
@app.route('/')
def home():
    return jsonify({'message': 'SkillSync AI Backend is running!', 'status': 'success'})

@app.route('/api/test')
def test():
    return jsonify({'message': 'Backend connection successful!', 'status': 'success'})

@app.route('/api/roles', methods=['GET'])
def get_roles():
    return jsonify({'roles': list(ROLE_SKILLS.keys())})

@app.route('/api/upload-resume', methods=['POST'])
def upload_resume():
    file = request.files.get('resume')
    if not file or not file.filename:
        return jsonify({'error': 'Resume file is required.'}), 400
    if not file.filename.lower().endswith(('.pdf', '.docx', '.txt')):
        return jsonify({'error': 'Only PDF, DOCX and TXT files are supported.'}), 400
    resume_text = extract_resume_text(file)
    if not resume_text:
        return jsonify({'error': 'Could not extract readable resume text.'}), 400
    return jsonify({'success': True, 'filename': file.filename, 'resume_text': resume_text,
                    'resumeText': resume_text, 'detected_skills': detect_skills(resume_text),
                    'message': 'Resume uploaded successfully.'})

@app.route('/api/analyze', methods=['POST'])
@app.route('/api/career-analysis', methods=['POST'])
def analyze():
    data = (request.get_json(silent=True) or {}) if request.is_json else request.form
    resume_text = first_value(data, 'resumeText', 'resume_text', default='')
    resume_file = request.files.get('resume')
    if resume_file and resume_file.filename:
        if not resume_file.filename.lower().endswith(('.pdf', '.docx', '.txt')):
            return jsonify({'error': 'Only PDF, DOCX and TXT files are supported.'}), 400
        resume_text = extract_resume_text(resume_file)
    role = valid_role(first_value(data, 'targetRole', 'target_role', default='Software Engineer'))
    github_username = first_value(data, 'githubUsername', 'github_username', default='')
    leetcode_username = first_value(data, 'leetcodeUsername', 'leetcode_username', default='')
    if not normalize_text(resume_text):
        return jsonify({'error': 'Could not extract readable resume text. Please upload a valid resume.'}), 400
    result = run_analysis(resume_text, role, github_username, leetcode_username)
    return jsonify({'success': True, **result})

@app.route('/api/skill-gap', methods=['POST'])
def skill_gap():
    data = request.get_json(silent=True) or {}
    role = valid_role(first_value(data, 'targetRole', 'target_role', default='Software Engineer'))
    resume_text = first_value(data, 'resumeText', 'resume_text', default='')
    analysis = data.get('analysis')
    if not resume_text and isinstance(analysis, dict):
        resume_text = first_value(analysis, 'resume_text', 'resumeText', default='')
    return jsonify(build_skill_gap(role, resume_text))

@app.route('/api/readiness-score', methods=['POST'])
def readiness_score():
    data = request.get_json(silent=True) or {}
    analysis = data.get('analysis') or {}
    resume = analysis.get('resume') or {}
    github = analysis.get('github') or {}
    leetcode = analysis.get('leetcode') or {}
    role_data = analysis.get('role_matching') or {}
    resume_score = safe_float(first_value(data, 'resumeScore', 'resume_score', default=resume.get('score', 0)))
    github_score = safe_float(first_value(data, 'githubScore', 'github_score', default=github.get('score', 0)))
    leetcode_score_value = safe_float(first_value(data, 'leetcodeScore', 'leetcode_score', default=leetcode.get('score', 0)))
    github_available = as_bool(first_value(data, 'githubAvailable', 'github_available', default=github.get('available', False)))
    leetcode_available = as_bool(first_value(data, 'leetcodeAvailable', 'leetcode_available', default=leetcode.get('available', False)))
    match_percentage = safe_float(first_value(data, 'roleMatchPercentage', 'role_match_percentage', default=role_data.get('match_percentage', 0)))
    return jsonify(calculate_readiness(resume_score, github_score, github_available,
                                       leetcode_score_value, leetcode_available, match_percentage))

@app.route('/api/recommendations', methods=['POST'])
def recommendations():
    data = request.get_json(silent=True) or {}
    return jsonify({'recommendations': generate_week1_recommendations(data.get('analysis') or {})})

@app.route('/api/what-if', methods=['POST'])
@app.route('/api/simulate', methods=['POST'])
@app.route('/api/what-if-simulator', methods=['POST'])
def simulate_scenario():
    data = request.get_json(silent=True) or {}
    analysis = data.get('analysis')
    if not isinstance(analysis, dict):
        return jsonify({'success': False, 'error': 'Analysis data is required.'}), 400
    try:
        result = simulate_what_if(analysis, data.get('scenario') or data.get('question') or '')
        return jsonify({'success': True, **result})
    except ValueError as error:
        return jsonify({'success': False, 'error': str(error)}), 400

# DEBUG ENDPOINTS
@app.route('/api/debug/github-auth', methods=['GET'])
def debug_github_auth():
    token = os.getenv('GITHUB_TOKEN', '').strip()
    try:
        response = requests.get(GITHUB_API + '/rate_limit', headers=github_headers(), timeout=REQUEST_TIMEOUT)
        try:
            data = response.json()
        except ValueError:
            data = {}
        core = (data.get('resources') or {}).get('core') or {} if isinstance(data, dict) else {}
        return jsonify({'token_detected': bool(token), 'github_status': response.status_code,
                        'limit': core.get('limit'), 'remaining': core.get('remaining'),
                        'reset': core.get('reset'), 'message': data.get('message') if isinstance(data, dict) else None})
    except requests.RequestException as error:
        return jsonify({'token_detected': bool(token), 'error': str(error)}), 502

@app.route('/api/debug/github/<username>')
def debug_github(username):
    return jsonify(analyze_github(username))

@app.route('/api/debug/leetcode/<username>')
def debug_leetcode(username):
    return jsonify(analyze_leetcode(username))

@app.route('/api/debug/cache/github')
def debug_github_cache():
    with GITHUB_CACHE_LOCK:
        cached_users = list(GITHUB_CACHE.keys())
    return jsonify({'cached_profiles': cached_users, 'cache_duration_minutes': 30, 'count': len(cached_users)})

@app.errorhandler(413)
def request_too_large(error):
    return jsonify({'error': 'Maximum upload size is 5 MB.'}), 413

@app.errorhandler(404)
def not_found(error):
    return jsonify({'error': 'API endpoint not found.'}), 404

@app.errorhandler(500)
def server_error(error):
    app.logger.error('Internal server error: %s', error)
    return jsonify({'error': 'Internal server error.'}), 500

if __name__ == '__main__':
    print('GitHub token detected:', bool(os.getenv('GITHUB_TOKEN', '').strip()))
    app.run(host='127.0.0.1', port=5000, debug=True)
